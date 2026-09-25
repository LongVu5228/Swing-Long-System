"""
One-off research script: test Qullamaggie's actual scanner mechanic (Section 7.2 of
KristjanDatabase/QULLAMAGGIE_TRADING_LESSONS.md -- "rank by relative strength across a
handful of lookback windows: 1/3/6 months") as a chart-pattern classifier, instead of
the 200MA slope/position test in analyze_chart_pattern_ma_trend.py.

Computes trailing 1M/3M/6M price change (pre-gap close vs. close ~21/~63/~126 trading
days earlier, no lookahead -- same pre_gap_close convention as the MA script) for every
human-labeled EP V5 event, then checks whether this cleanly separates DT-family
(DT/DT SW/DT U) from UT-family (UT/UTU), same comparison as the MA analysis so the two
approaches are directly comparable.

Usage:
    python analyze_chart_pattern_relative_strength.py               # full run
    python analyze_chart_pattern_relative_strength.py --smoke-test   # first 20 events, writes _SMOKETEST.csv instead
"""
import argparse
import os
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date

import openpyxl
import pandas as pd
from dateutil.relativedelta import relativedelta
from dotenv import load_dotenv
from tqdm import tqdm

load_dotenv(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".env"))

SCRIPTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "Scripts")
sys.path.insert(0, SCRIPTS_DIR)

from fill_episodic_pivots_v3 import get_daily_bars, PLAN_CUTOFF  # noqa: E402
from build_benzinga_candidate_list import resolve_historical_ticker  # noqa: E402
from build_v2_features import clean_ticker  # noqa: E402

EP_XLSX = os.path.join(os.path.dirname(__file__), "..", "EP V5.xlsx")
SHEET_NAME = "Data"
OUTPUT_CSV = os.path.join(os.path.dirname(__file__), "chart_pattern_relative_strength.csv")
SMOKETEST_CSV = os.path.join(os.path.dirname(__file__), "chart_pattern_relative_strength_SMOKETEST.csv")
# Trading-day approximations for 1/3/6 calendar months, matching Qullamaggie's own
# "one month, three months, six months" scan windows (Section 7.2).
LOOKBACKS = {"1m": 21, "3m": 63, "6m": 126}
TODAY = date.today()


def load_events():
    wb = openpyxl.load_workbook(EP_XLSX, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        ticker = row[idx["ticker"]]
        rdate = row[idx["reaction_date"]]
        pattern = row[idx["Chart Pattern"]]
        if not ticker or not rdate or not pattern or pattern == "DELISTED":
            continue
        event_date = rdate.date() if hasattr(rdate, "date") else date.fromisoformat(str(rdate))
        events.append({"ticker": ticker, "event_date": event_date, "pattern": pattern})
    return events


def analyze_event(raw_ticker: str, event_date: date, pattern: str) -> dict:
    ticker = clean_ticker(raw_ticker)
    lo = max(event_date - relativedelta(months=9), PLAN_CUTOFF)
    hi = event_date
    if lo > hi:
        return {"ticker": raw_ticker, "date": event_date.isoformat(), "pattern": pattern,
                "status": "before_polygon_plan_cutoff"}

    bars = get_daily_bars(ticker, lo.isoformat(), hi.isoformat())
    if bars is None:
        alt = resolve_historical_ticker(ticker, event_date)
        if alt != ticker:
            bars = get_daily_bars(alt, lo.isoformat(), hi.isoformat())
    if bars is None:
        return {"ticker": raw_ticker, "date": event_date.isoformat(), "pattern": pattern, "status": "no_daily_bars"}

    prior = bars[bars["date"] < event_date].reset_index(drop=True)
    max_needed = max(LOOKBACKS.values())
    if len(prior) < max_needed + 1:
        return {"ticker": raw_ticker, "date": event_date.isoformat(), "pattern": pattern,
                "status": "insufficient_history"}

    closes = prior["close"]
    pre_gap_close = closes.iloc[-1]

    out = {"ticker": raw_ticker, "date": event_date.isoformat(), "pattern": pattern, "status": "OK"}
    for label, n_days in LOOKBACKS.items():
        ref_close = closes.iloc[-1 - n_days]
        out[f"rs_{label}_pct"] = (pre_gap_close / ref_close - 1) * 100
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--smoke-test", action="store_true")
    parser.add_argument("--workers", type=int, default=8)
    args = parser.parse_args()

    events = load_events()
    if args.smoke_test:
        events = events[:20]
    print(f"{len(events)} labeled events loaded from EP V5.xlsx")

    results = []
    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        futures = {
            executor.submit(analyze_event, e["ticker"], e["event_date"], e["pattern"]): e
            for e in events
        }
        exception_count = 0
        with tqdm(total=len(futures), desc="Relative strength analysis", unit="event") as pbar:
            for fut in as_completed(futures):
                e = futures[fut]
                try:
                    r = fut.result()
                except Exception as ex:
                    exception_count += 1
                    r = {"ticker": e["ticker"], "date": e["event_date"].isoformat(), "pattern": e["pattern"],
                         "status": f"exception: {type(ex).__name__}: {ex}"}
                results.append(r)
                pbar.update(1)
        if exception_count:
            print(f"\n{exception_count} events hit an unexpected exception (recorded in 'status').")

    df = pd.DataFrame(results)
    out_path = SMOKETEST_CSV if args.smoke_test else OUTPUT_CSV
    df.to_csv(out_path, index=False)
    print(f"\nWrote {len(df)} rows to {out_path}")
    print(f"status counts: {df['status'].value_counts().to_dict()}")

    ok = df[df["status"] == "OK"].copy()
    print(f"\n{len(ok)} events with complete 6M-lookback history\n")

    for label in LOOKBACKS:
        col = f"rs_{label}_pct"
        ok[f"{label}_negative"] = ok[col] < 0

    agg = {f"rs_{l}_pct": "median" for l in LOOKBACKS}
    agg.update({f"{l}_negative": lambda s: (s).mean() * 100 for l in LOOKBACKS})
    # named aggregation to avoid duplicate-lambda-name collisions
    named = {"n": ("pattern", "size")}
    for l in LOOKBACKS:
        named[f"median_rs_{l}_pct"] = (f"rs_{l}_pct", "median")
        named[f"pct_{l}_negative"] = (f"rs_{l}_pct", lambda s: (s < 0).mean() * 100)

    summary = ok.groupby("pattern").agg(**named).sort_values("pct_1m_negative", ascending=False)
    pd.set_option("display.width", 160)
    print(summary.round(1).to_string())


if __name__ == "__main__":
    main()
