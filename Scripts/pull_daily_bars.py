"""
Raw data layer: pull daily OHLCV from Polygon for every ticker in the EP candidate
universe and cache it locally as Parquet, one file per ticker.

This is the immutable/raw layer in the pipeline -- point-in-time feature engineering,
gap validation, and forward-return calculations all read from this cache rather than
re-hitting the API. Re-run to refresh/extend the cache; existing per-ticker files are
skipped unless --refresh is passed.
"""

import argparse
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pandas as pd
import requests
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from tqdm import tqdm

load_dotenv()

POLYGON_API_KEY = os.environ.get("POLYGON_API_KEY")
if not POLYGON_API_KEY:
    sys.exit("POLYGON_API_KEY not found. Check your .env file.")

INPUT_XLSX = os.path.join("Files", "EP", "All EP Scan.xlsx")
INPUT_SHEET = "OG No Dupes"
CACHE_DIR = os.path.join("Data", "daily_bars")
MANIFEST_PATH = os.path.join("Data", "daily_bars_manifest.csv")

BASE_URL = "https://api.polygon.io"
MAX_RETRIES = 5
MAX_WORKERS = 20
BUFFER_BEFORE_DAYS = 400  # covers ~252 trading days lookback (200DMA, 1y return) + weekends/holidays
BUFFER_AFTER_DAYS = 60    # covers ~20-40 trading day forward window (MFE/MAE, forward returns)

SESSION = requests.Session()
_adapter = HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS)
SESSION.mount("https://", _adapter)


def api_get(path, params=None):
    params = dict(params or {})
    params["apiKey"] = POLYGON_API_KEY
    url = f"{BASE_URL}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        resp = SESSION.get(url, params=params, timeout=30)
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code == 404:
            return None
        resp.raise_for_status()
        return resp.json()
    return None


def clean_ticker(raw_ticker):
    return raw_ticker.split("^")[0].strip()


def load_universe():
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[INPUT_SHEET]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        ticker = row[idx["Ticker"]]
        event_date = row[idx["Full Date"]]
        if not ticker or not event_date:
            continue
        events.append((clean_ticker(ticker), event_date))

    ranges = {}
    for ticker, event_date in events:
        d = date.fromisoformat(str(event_date))
        lo, hi = ranges.get(ticker, (d, d))
        ranges[ticker] = (min(lo, d), max(hi, d))
    return ranges


def get_daily_bars(ticker, from_date, to_date):
    path = f"/v2/aggs/ticker/{ticker}/range/1/day/{from_date}/{to_date}"
    data = api_get(path, {"adjusted": "true", "sort": "asc", "limit": 50000})
    if data is None:
        return None, "error"
    if not data.get("results"):
        return [], data.get("status", "no_results")
    return data["results"], data.get("status", "OK")


def pull_ticker(ticker, event_min, event_max, refresh):
    out_path = os.path.join(CACHE_DIR, f"{ticker}.parquet")
    from_date = (event_min - timedelta(days=BUFFER_BEFORE_DAYS)).isoformat()
    to_date = min(event_max + timedelta(days=BUFFER_AFTER_DAYS), date.today()).isoformat()

    if not refresh and os.path.exists(out_path):
        return {"ticker": ticker, "status": "cached", "from_date": from_date, "to_date": to_date,
                "bars": None, "first_bar": None, "last_bar": None}

    bars, status = get_daily_bars(ticker, from_date, to_date)
    if bars is None:
        return {"ticker": ticker, "status": "error", "from_date": from_date, "to_date": to_date,
                "bars": 0, "first_bar": None, "last_bar": None}
    if not bars:
        return {"ticker": ticker, "status": "no_data", "from_date": from_date, "to_date": to_date,
                "bars": 0, "first_bar": None, "last_bar": None}

    df = pd.DataFrame(bars).rename(columns={
        "t": "timestamp_ms", "o": "open", "h": "high", "l": "low",
        "c": "close", "v": "volume", "vw": "vwap", "n": "transactions",
    })
    df["date"] = df["timestamp_ms"].apply(
        lambda ms: datetime.fromtimestamp(ms / 1000, tz=timezone.utc).date().isoformat()
    )
    cols = ["date", "open", "high", "low", "close", "volume", "vwap", "transactions"]
    df = df[cols]

    os.makedirs(CACHE_DIR, exist_ok=True)
    df.to_parquet(out_path, index=False)

    return {"ticker": ticker, "status": "ok", "from_date": from_date, "to_date": to_date,
            "bars": len(df), "first_bar": df["date"].iloc[0], "last_bar": df["date"].iloc[-1]}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh", action="store_true", help="Re-pull tickers already cached")
    parser.add_argument("--limit", type=int, default=None, help="Only process the first N tickers (testing)")
    args = parser.parse_args()

    ranges = load_universe()
    tickers = sorted(ranges.items())
    if args.limit:
        tickers = tickers[: args.limit]

    print(f"{len(tickers)} unique tickers to process ({'refresh' if args.refresh else 'skip cached'})")

    results = []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {
            executor.submit(pull_ticker, ticker, lo, hi, args.refresh): ticker
            for ticker, (lo, hi) in tickers
        }
        with tqdm(total=len(futures), desc="Pulling daily bars", unit="ticker") as pbar:
            for future in as_completed(futures):
                results.append(future.result())
                pbar.update(1)

    os.makedirs("Data", exist_ok=True)
    pd.DataFrame(results).to_csv(MANIFEST_PATH, index=False)

    status_counts = pd.DataFrame(results)["status"].value_counts()
    print("\nStatus breakdown:")
    print(status_counts.to_string())
    print(f"\nManifest written to {MANIFEST_PATH}")
    print(f"Cache directory: {CACHE_DIR}")


if __name__ == "__main__":
    main()
