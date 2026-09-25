"""
Second pass over the Benzinga Earnings Beat-Miss output: for rows still N/A, resolve
the ticker's historical symbol via Polygon's ticker-events history and re-query
Benzinga earnings under that symbol. Only overwrites rows that come back with a match.
"""

import os
import sys
from datetime import date

import openpyxl
import pandas as pd
from tqdm import tqdm

sys.path.insert(0, os.path.dirname(__file__))
from build_v2_features import clean_ticker, resolve_historical_ticker
from pull_benzinga_earnings_to_file import get_earnings_history, find_match, yoy_growth, beat_miss, na

OUT_PATH = os.path.join("Files", "EP", "Data Pulls", "Benzinga Earnings Beat-Miss.xlsx")


def main():
    wb = openpyxl.load_workbook(OUT_PATH, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}

    targets = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[idx["Fiscal Quarter"]] != "N/A":
            continue
        targets.append({
            "row": row_num, "ticker_raw": row[idx["Ticker"]],
            "event_date": date.fromisoformat(row[idx["Full Date"]]),
        })

    print(f"{len(targets)} N/A rows eligible for rename resolution")

    recovered = {}
    for t in tqdm(targets, desc="Resolving renamed tickers", unit="row"):
        ticker = clean_ticker(t["ticker_raw"])
        resolved = resolve_historical_ticker(ticker, t["event_date"])
        if resolved == ticker:
            continue
        records = get_earnings_history(resolved)
        m = find_match(records, t["event_date"])
        if m:
            recovered[t["row"]] = (resolved, m)

    print(f"\nRecovered {len(recovered)} / {len(targets)} rows")
    for row_num, (resolved, m) in list(recovered.items())[:30]:
        orig = next(t["ticker_raw"] for t in targets if t["row"] == row_num)
        print(f"  row {row_num}: {orig} -> {resolved}")

    if not recovered:
        return

    wb2 = openpyxl.load_workbook(OUT_PATH, data_only=False)
    ws2 = wb2.active
    col = {name: i + 1 for i, name in enumerate(header)}

    for row_num, (resolved, m) in recovered.items():
        ws2.cell(row=row_num, column=col["Fiscal Quarter"], value=f"{m.get('fiscal_period')} {m.get('fiscal_year')}")
        ws2.cell(row=row_num, column=col["Report Date"], value=m["date"].isoformat())
        ws2.cell(row=row_num, column=col["Actual EPS"], value=na(m.get("actual_eps")))
        ws2.cell(row=row_num, column=col["Estimated EPS"], value=na(m.get("estimated_eps")))
        eps_surp_pct = m.get("eps_surprise_percent") * 100 if m.get("eps_surprise_percent") is not None else None
        ws2.cell(row=row_num, column=col["EPS Surprise %"], value=na(eps_surp_pct))
        ws2.cell(row=row_num, column=col["EPS Beat?"], value=beat_miss(m.get("eps_surprise")))
        ws2.cell(row=row_num, column=col["Actual Revenue"], value=na(m.get("actual_revenue")))
        ws2.cell(row=row_num, column=col["Estimated Revenue"], value=na(m.get("estimated_revenue")))
        rev_surp_pct = m.get("revenue_surprise_percent") * 100 if m.get("revenue_surprise_percent") is not None else None
        ws2.cell(row=row_num, column=col["Revenue Surprise %"], value=na(rev_surp_pct))
        ws2.cell(row=row_num, column=col["Revenue Beat?"], value=beat_miss(m.get("revenue_surprise")))
        ws2.cell(row=row_num, column=col["Previous Year EPS"], value=na(m.get("previous_eps")))
        ws2.cell(row=row_num, column=col["Previous Year Revenue"], value=na(m.get("previous_revenue")))
        ws2.cell(row=row_num, column=col["EPS YoY Growth %"], value=na(yoy_growth(m.get("actual_eps"), m.get("previous_eps"))))
        ws2.cell(row=row_num, column=col["Revenue YoY Growth %"], value=na(yoy_growth(m.get("actual_revenue"), m.get("previous_revenue"))))

    wb2.save(OUT_PATH)
    print(f"\nWrote {len(recovered)} recovered rows to {OUT_PATH}")


if __name__ == "__main__":
    main()
