"""
Pull IPO date, pre-gap volume baselines, gap-day totals, and opening-day minute-candle
stats (1/5/10/15/30/60 min) into the "Pasted Data" tab of EP Scan V2.xlsx.

New columns:
  IPO Date
  Pre-Gap 30D Avg Dollar Volume, Pre-Gap 30D Avg Share Volume
  Pre-Gap 100D Avg Dollar Volume, Pre-Gap 100D Avg Share Volume
  Gap Day Total Volume, Gap Day Total Dollar Volume
  {1,5,10,15,30,60}M Candle Close / Volume / Dollar Volume

All "pre-gap" stats use daily bars strictly BEFORE the event date (point-in-time,
no lookahead). Candle stats are built from 1-minute bars in the regular session
(9:30am ET open), bucketed by wall-clock time so sparse/no-trade minutes don't
shift the window -- e.g. the "5M" candle is bars with 9:30 <= t < 9:35 ET, whatever
bars actually exist in that span, not just "the first 5 bars returned".
"""

import argparse
import os
import sys
import time
import zoneinfo
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pandas as pd
import requests
from openpyxl.utils.datetime import from_excel
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from tqdm import tqdm

sys.path.insert(0, os.path.dirname(__file__))
from build_v2_features import clean_ticker, resolve_historical_ticker, PLAN_CUTOFF

load_dotenv()
POLYGON_API_KEY = os.environ.get("POLYGON_API_KEY")
if not POLYGON_API_KEY:
    sys.exit("POLYGON_API_KEY not found. Check your .env file.")

INPUT_XLSX = os.path.join("Files", "EP", "EP Scan V2.xlsx")
SHEET_NAME = "Pasted Data"

BASE_URL = "https://api.polygon.io"
MAX_RETRIES = 5
MAX_WORKERS = 15
ET = zoneinfo.ZoneInfo("America/New_York")
CANDLE_WINDOWS = [1, 5, 10, 15, 30, 60]
TODAY = date.today()

SESSION = requests.Session()
SESSION.mount("https://", HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS))

_ipo_cache = {}
_ipo_lock_keys = set()


def api_get(path, params=None):
    params = dict(params or {})
    params["apiKey"] = POLYGON_API_KEY
    url = f"{BASE_URL}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        resp = SESSION.get(url, params=params, timeout=30)
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code in (403, 404):
            return None
        resp.raise_for_status()
        return resp.json()
    return None


def get_daily_bars(ticker, from_date, to_date):
    path = f"/v2/aggs/ticker/{ticker}/range/1/day/{from_date}/{to_date}"
    data = api_get(path, {"adjusted": "true", "sort": "asc", "limit": 50000})
    if data is None or not data.get("results"):
        return None
    df = pd.DataFrame(data["results"])
    df["date"] = df["t"].apply(lambda ms: datetime.fromtimestamp(ms / 1000, tz=timezone.utc).date())
    return df[["date", "v", "vw", "c"]].rename(columns={"v": "volume", "vw": "vwap", "c": "close"})


def get_minute_bars(ticker, event_date):
    d = event_date.isoformat()
    path = f"/v2/aggs/ticker/{ticker}/range/1/minute/{d}/{d}"
    data = api_get(path, {"adjusted": "true", "sort": "asc", "limit": 1000})
    if data is None or not data.get("results"):
        return None
    df = pd.DataFrame(data["results"])
    df["dt_et"] = df["t"].apply(lambda ms: datetime.fromtimestamp(ms / 1000, tz=timezone.utc).astimezone(ET))
    return df[["dt_et", "o", "h", "l", "c", "v", "vw"]]


def get_ipo_date(ticker):
    if ticker in _ipo_cache:
        return _ipo_cache[ticker]
    data = api_get(f"/v3/reference/tickers/{ticker}")
    list_date = None
    if data and data.get("results", {}).get("list_date"):
        list_date = data["results"]["list_date"]
    else:
        search = api_get("/v3/reference/tickers", {"ticker": ticker, "active": "false", "limit": 1})
        if search and search.get("results"):
            delisted_utc = search["results"][0].get("delisted_utc")
            if delisted_utc:
                as_of = (datetime.fromisoformat(delisted_utc.replace("Z", "+00:00")).date()
                         - timedelta(days=30)).isoformat()
                data2 = api_get(f"/v3/reference/tickers/{ticker}", {"date": as_of})
                if data2 and data2.get("results", {}).get("list_date"):
                    list_date = data2["results"]["list_date"]
    _ipo_cache[ticker] = list_date
    return list_date


def compute_candles(minute_df, event_date):
    out = {}
    if minute_df is None or minute_df.empty:
        for m in CANDLE_WINDOWS:
            out[f"close_{m}m"], out[f"vol_{m}m"], out[f"dvol_{m}m"] = None, None, None
        return out

    session_open = datetime(event_date.year, event_date.month, event_date.day, 9, 30, tzinfo=ET)
    regular = minute_df[minute_df["dt_et"] >= session_open].sort_values("dt_et")

    for m in CANDLE_WINDOWS:
        window_end = session_open + timedelta(minutes=m)
        window = regular[regular["dt_et"] < window_end]
        if window.empty:
            out[f"close_{m}m"], out[f"vol_{m}m"], out[f"dvol_{m}m"] = None, None, None
        else:
            out[f"close_{m}m"] = window["c"].iloc[-1]
            out[f"vol_{m}m"] = window["v"].sum()
            out[f"dvol_{m}m"] = (window["v"] * window["vw"]).sum()
    return out


def na(v):
    return v if v is not None else "N/A"


def process_ticker(ticker, ticker_events):
    lo = min(e["event_date"] for e in ticker_events) - timedelta(days=170)
    hi = max(e["event_date"] for e in ticker_events)
    lo = max(lo, PLAN_CUTOFF)

    daily = get_daily_bars(ticker, lo.isoformat(), hi.isoformat()) if lo <= hi else None
    resolved_ticker = ticker
    if daily is None and lo <= hi:
        resolved = resolve_historical_ticker(ticker, min(e["event_date"] for e in ticker_events))
        if resolved != ticker:
            daily = get_daily_bars(resolved, lo.isoformat(), hi.isoformat())
            resolved_ticker = resolved

    ipo_date = get_ipo_date(resolved_ticker) if daily is not None else None

    results = {}
    for e in ticker_events:
        row = {"ipo_date": ipo_date}
        if daily is None:
            row.update({"dvol30": None, "svol30": None, "dvol100": None, "svol100": None,
                        "gap_vol": None, "gap_dvol": None})
        else:
            prior = daily[daily["date"] < e["event_date"]].sort_values("date")
            prior30 = prior.tail(30)
            prior100 = prior.tail(100)
            # Matches the user's live TradingView/P123 convention: pre-gap close * average
            # DAILY SHARE volume -- NOT mean(daily volume * that day's own price), which
            # understates this when the stock trended up over the window (and vice versa).
            pre_gap_close = prior["close"].iloc[-1] if not prior.empty else None
            row["svol30"] = prior30["volume"].mean() if not prior30.empty else None
            row["dvol30"] = pre_gap_close * row["svol30"] if pre_gap_close is not None and row["svol30"] is not None else None
            row["svol100"] = prior100["volume"].mean() if not prior100.empty else None
            row["dvol100"] = pre_gap_close * row["svol100"] if pre_gap_close is not None and row["svol100"] is not None else None

            gap_bar = daily[daily["date"] == e["event_date"]]
            if not gap_bar.empty:
                row["gap_vol"] = gap_bar["volume"].iloc[0]
                row["gap_dvol"] = (gap_bar["volume"] * gap_bar["vwap"]).iloc[0]
            else:
                row["gap_vol"], row["gap_dvol"] = None, None

        minute_bars = get_minute_bars(resolved_ticker, e["event_date"])
        row.update(compute_candles(minute_bars, e["event_date"]))
        results[e["row"]] = row
    return results


def load_events(limit=None):
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ticker_raw = row[idx["Ticker"]]
        fd = row[idx["Full Date"]]
        if not ticker_raw or fd is None:
            continue
        if isinstance(fd, datetime):
            event_date = fd.date()
        elif isinstance(fd, (int, float)):
            event_date = from_excel(fd).date()
        else:
            event_date = date.fromisoformat(str(fd))
        events.append({"row": row_num, "ticker_raw": ticker_raw, "ticker": clean_ticker(ticker_raw),
                        "event_date": event_date})
    if limit:
        events = events[:limit]
    return events


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    events = load_events(limit=args.limit)
    by_ticker = {}
    for e in events:
        by_ticker.setdefault(e["ticker"], []).append(e)
    print(f"{len(events)} event rows across {len(by_ticker)} unique tickers")

    all_results = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(process_ticker, t, evs): t for t, evs in by_ticker.items()}
        with tqdm(total=len(futures), desc="Building intraday features", unit="ticker") as pbar:
            for future in as_completed(futures):
                all_results.update(future.result())
                pbar.update(1)

    if args.dry_run:
        rows = []
        for e in events:
            r = all_results.get(e["row"], {})
            rows.append({"ticker": e["ticker"], "event_date": e["event_date"].isoformat(),
                         "ipo": r.get("ipo_date"), "dvol30": r.get("dvol30"), "svol100": r.get("svol100"),
                         "gap_vol": r.get("gap_vol"), "gap_dvol": r.get("gap_dvol"),
                         "close_1m": r.get("close_1m"), "vol_1m": r.get("vol_1m"), "dvol_1m": r.get("dvol_1m"),
                         "close_5m": r.get("close_5m"), "vol_5m": r.get("vol_5m"),
                         "close_60m": r.get("close_60m"), "vol_60m": r.get("vol_60m")})
        print(pd.DataFrame(rows).to_string(index=False))
        return

    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=False)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    next_col = len(header) + 1

    columns = ["IPO Date",
               "Pre-Gap 30D Avg Dollar Volume", "Pre-Gap 30D Avg Share Volume",
               "Pre-Gap 100D Avg Dollar Volume", "Pre-Gap 100D Avg Share Volume",
               "Gap Day Total Volume", "Gap Day Total Dollar Volume"]
    for m in CANDLE_WINDOWS:
        columns += [f"{m}M Candle Close", f"{m}M Candle Volume", f"{m}M Candle Dollar Volume"]

    col_map = {}
    for name in columns:
        ws.cell(row=1, column=next_col, value=name)
        col_map[name] = next_col
        next_col += 1

    for e in events:
        r = all_results.get(e["row"], {})
        row_num = e["row"]
        ws.cell(row=row_num, column=col_map["IPO Date"], value=na(r.get("ipo_date")))
        ws.cell(row=row_num, column=col_map["Pre-Gap 30D Avg Dollar Volume"], value=na(r.get("dvol30")))
        ws.cell(row=row_num, column=col_map["Pre-Gap 30D Avg Share Volume"], value=na(r.get("svol30")))
        ws.cell(row=row_num, column=col_map["Pre-Gap 100D Avg Dollar Volume"], value=na(r.get("dvol100")))
        ws.cell(row=row_num, column=col_map["Pre-Gap 100D Avg Share Volume"], value=na(r.get("svol100")))
        ws.cell(row=row_num, column=col_map["Gap Day Total Volume"], value=na(r.get("gap_vol")))
        ws.cell(row=row_num, column=col_map["Gap Day Total Dollar Volume"], value=na(r.get("gap_dvol")))
        for m in CANDLE_WINDOWS:
            ws.cell(row=row_num, column=col_map[f"{m}M Candle Close"], value=na(r.get(f"close_{m}m")))
            ws.cell(row=row_num, column=col_map[f"{m}M Candle Volume"], value=na(r.get(f"vol_{m}m")))
            ws.cell(row=row_num, column=col_map[f"{m}M Candle Dollar Volume"], value=na(r.get(f"dvol_{m}m")))

    wb.save(INPUT_XLSX)
    print(f"\nWrote {len(events)} rows / {len(columns)} new columns to {INPUT_XLSX}")


if __name__ == "__main__":
    main()
