"""
Second pass over EP Scan V2.xlsx: for rows where Gap Day Open is still N/A and the
event date is within Polygon's reachable window (not the plan-cutoff case), resolve
the ticker's actual historical symbol via Polygon's ticker-events history and
recompute all the feature columns (gap OHLC, prior ATH, 1/3/6M high/close) under
that symbol. Only overwrites rows that come back with real data.
"""

import argparse
from datetime import date, datetime

import openpyxl
import pandas as pd
from dateutil.relativedelta import relativedelta
from tqdm import tqdm

from build_v2_features import (
    INPUT_XLSX, SHEET_NAME, PLAN_CUTOFF, WINDOWS_MONTHS,
    clean_ticker, resolve_historical_ticker, get_daily_bars, compute_row, na,
)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}

    targets = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[idx["Gap Day Open"]] != "N/A":
            continue
        fd = row[idx["Full Date"]]
        event_date = fd.date() if isinstance(fd, datetime) else date.fromisoformat(str(fd))
        if event_date < PLAN_CUTOFF:
            continue
        targets.append({"row": row_num, "ticker_raw": row[idx["Ticker"]], "event_date": event_date})

    print(f"{len(targets)} N/A rows eligible for rename resolution")

    recovered = {}
    for t in tqdm(targets, desc="Resolving renamed tickers", unit="row"):
        ticker = clean_ticker(t["ticker_raw"])
        resolved = resolve_historical_ticker(ticker, t["event_date"])
        lo = PLAN_CUTOFF
        hi = min(date.today(), t["event_date"] + relativedelta(months=max(WINDOWS_MONTHS)))
        bars = get_daily_bars(resolved, lo.isoformat(), hi.isoformat())
        if bars is None:
            continue
        r = compute_row(bars, t["event_date"])
        if r["gap_open"] is None:
            continue
        r["resolved_ticker"] = resolved
        recovered[t["row"]] = r

    print(f"\nRecovered {len(recovered)} / {len(targets)} rows")
    for row_num, r in list(recovered.items())[:30]:
        orig = next(t["ticker_raw"] for t in targets if t["row"] == row_num)
        print(f"  row {row_num}: {orig} -> {r['resolved_ticker']}")

    if args.dry_run or not recovered:
        return

    wb2 = openpyxl.load_workbook(INPUT_XLSX, data_only=False)
    ws2 = wb2[SHEET_NAME]
    col_map = {name: i + 1 for i, name in enumerate(header)}

    for row_num, r in recovered.items():
        ws2.cell(row=row_num, column=col_map["Gap Day Open"], value=na(r["gap_open"]))
        ws2.cell(row=row_num, column=col_map["Gap Day Close"], value=na(r["gap_close"]))
        ws2.cell(row=row_num, column=col_map["Prior ATH Price"], value=na(r["ath_price"]))
        ws2.cell(row=row_num, column=col_map["Prior ATH Date"], value=na(r["ath_date"]))
        ws2.cell(row=row_num, column=col_map["Prior ATH Data Truncated"], value=r["ath_truncated"])
        for m in WINDOWS_MONTHS:
            ws2.cell(row=row_num, column=col_map[f"{m}M High"], value=na(r.get(f"high_{m}m")))
            ws2.cell(row=row_num, column=col_map[f"{m}M High Date"], value=na(r.get(f"high_{m}m_date")))
            ws2.cell(row=row_num, column=col_map[f"{m}M Window Complete"], value=r.get(f"{m}m_complete"))
            ws2.cell(row=row_num, column=col_map[f"{m}M Plan Data Truncated"], value=r.get(f"{m}m_truncated"))
            ws2.cell(row=row_num, column=col_map[f"{m}M Close"], value=na(r.get(f"close_{m}m")))
            ws2.cell(row=row_num, column=col_map[f"{m}M Close Date"], value=na(r.get(f"close_{m}m_date")))

    wb2.save(INPUT_XLSX)
    print(f"\nWrote {len(recovered)} recovered rows to {INPUT_XLSX}")


if __name__ == "__main__":
    main()
