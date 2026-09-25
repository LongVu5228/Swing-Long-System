"""
Pull EPS/Revenue Beat-Miss + YoY growth from Benzinga's earnings feed (via Polygon/
Massive's /benzinga/v1/earnings) for every event in Master Order of Tickers.xlsx.

Unlike the SEC-financials approach (pull_earnings_growth_to_file.py), this has real
analyst consensus estimates and surprise% built in, plus previous_eps/previous_revenue
(same period prior year) for YoY growth -- no manual quarter-matching heuristics needed.

Matches by report date proximity: Benzinga's `date` field is the actual earnings
report date, which should sit right at (or within a couple days of) our EP event date,
since these EP events were identified as earnings-driven gaps in the first place.

Output goes to a standalone file in Data Pulls/, row-aligned to Master Order of
Tickers -- not written into any working file directly.
"""

import argparse
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
import requests
from openpyxl.utils.datetime import from_excel
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from tqdm import tqdm

sys.path.insert(0, os.path.dirname(__file__))
from build_v2_features import clean_ticker, resolve_historical_ticker

load_dotenv()
POLYGON_API_KEY = os.environ.get("POLYGON_API_KEY")
if not POLYGON_API_KEY:
    sys.exit("POLYGON_API_KEY not found. Check your .env file.")

MASTER_ORDER_XLSX = os.path.join("Files", "EP", "Master Order of tickers", "Master Order of Tickers.xlsx")
OUTPUT_DIR = os.path.join("Files", "EP", "Data Pulls")

BASE_URL = "https://api.polygon.io"
MAX_RETRIES = 5
MAX_WORKERS = 10
MATCH_TOLERANCE_DAYS = 5  # event date vs Benzinga's report date

SESSION = requests.Session()
SESSION.mount("https://", HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS))


def api_get(path, params=None):
    params = dict(params or {})
    params["apiKey"] = POLYGON_API_KEY
    url = f"{BASE_URL}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            resp = SESSION.get(url, params=params, timeout=30)
        except requests.exceptions.RequestException:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code in (403, 404):
            return None
        resp.raise_for_status()
        return resp.json()
    return None


def get_earnings_history(ticker, limit=40):
    data = api_get("/benzinga/v1/earnings", {
        "ticker": ticker, "date_status": "confirmed", "limit": limit, "sort": "date.desc",
    })
    if not data or not data.get("results"):
        return []
    out = []
    for r in data["results"]:
        d = r.get("date")
        if not d:
            continue
        out.append({**r, "date": date.fromisoformat(d)})
    return out


def find_match(records, event_date):
    if not records:
        return None
    closest = min(records, key=lambda r: abs((r["date"] - event_date).days))
    if abs((closest["date"] - event_date).days) > MATCH_TOLERANCE_DAYS:
        return None
    return closest


def yoy_growth(actual, previous):
    if actual is None or previous in (None, 0):
        return None
    return (actual - previous) / abs(previous) * 100


def beat_miss(surprise):
    if surprise is None:
        return "N/A"
    return "Y" if surprise > 0 else "N"


def process_ticker(ticker, ticker_events):
    records = get_earnings_history(ticker)
    resolved = ticker
    if not records:
        alt = resolve_historical_ticker(ticker, min(e["event_date"] for e in ticker_events))
        if alt != ticker:
            records = get_earnings_history(alt)
            resolved = alt

    results = {}
    for e in ticker_events:
        m = find_match(records, e["event_date"])
        results[e["row"]] = {"resolved_ticker": resolved, "match": m}
    return results


def na(v):
    return v if v is not None else "N/A"


def load_master_order(limit=None):
    wb = openpyxl.load_workbook(MASTER_ORDER_XLSX, data_only=True)
    ws = wb["Sheet1"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ticker_raw = row[idx["Ticker"]]
        fd = row[idx["Full Date"]]
        if not ticker_raw or fd is None:
            continue
        if isinstance(fd, datetime):
            event_date = fd.date()
        elif isinstance(fd, (int, float)):
            event_date = from_excel(fd).date()
        else:
            event_date = date.fromisoformat(str(fd))
        events.append({"row": row_num, "ticker_raw": ticker_raw, "ticker": clean_ticker(ticker_raw),
                        "event_date": event_date})
    if limit:
        events = events[:limit]
    return events


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--label", type=str, default=None)
    args = parser.parse_args()

    all_events = load_master_order()
    events = all_events[args.offset: args.offset + args.limit] if args.limit else all_events
    by_ticker = {}
    for e in events:
        by_ticker.setdefault(e["ticker"], []).append(e)
    print(f"{len(events)} event rows across {len(by_ticker)} unique tickers")

    all_results = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(process_ticker, t, evs): t for t, evs in by_ticker.items()}
        with tqdm(total=len(futures), desc="Pulling Benzinga earnings", unit="ticker") as pbar:
            for future in as_completed(futures):
                all_results.update(future.result())
                pbar.update(1)

    out_rows = []
    matched = 0
    for e in events:
        r = all_results.get(e["row"], {})
        m = r.get("match")
        if m:
            matched += 1
            row = {
                "Full Date": e["event_date"].isoformat(), "Ticker": e["ticker_raw"],
                "Fiscal Quarter": f"{m.get('fiscal_period')} {m.get('fiscal_year')}",
                "Report Date": m["date"].isoformat(),
                "Actual EPS": na(m.get("actual_eps")), "Estimated EPS": na(m.get("estimated_eps")),
                "EPS Surprise %": na(m.get("eps_surprise_percent") * 100 if m.get("eps_surprise_percent") is not None else None),
                "EPS Beat?": beat_miss(m.get("eps_surprise")),
                "Actual Revenue": na(m.get("actual_revenue")), "Estimated Revenue": na(m.get("estimated_revenue")),
                "Revenue Surprise %": na(m.get("revenue_surprise_percent") * 100 if m.get("revenue_surprise_percent") is not None else None),
                "Revenue Beat?": beat_miss(m.get("revenue_surprise")),
                "Previous Year EPS": na(m.get("previous_eps")), "Previous Year Revenue": na(m.get("previous_revenue")),
                "EPS YoY Growth %": na(yoy_growth(m.get("actual_eps"), m.get("previous_eps"))),
                "Revenue YoY Growth %": na(yoy_growth(m.get("actual_revenue"), m.get("previous_revenue"))),
            }
        else:
            row = {
                "Full Date": e["event_date"].isoformat(), "Ticker": e["ticker_raw"],
                "Fiscal Quarter": "N/A", "Report Date": "N/A",
                "Actual EPS": "N/A", "Estimated EPS": "N/A", "EPS Surprise %": "N/A", "EPS Beat?": "N/A",
                "Actual Revenue": "N/A", "Estimated Revenue": "N/A", "Revenue Surprise %": "N/A", "Revenue Beat?": "N/A",
                "Previous Year EPS": "N/A", "Previous Year Revenue": "N/A",
                "EPS YoY Growth %": "N/A", "Revenue YoY Growth %": "N/A",
            }
        out_rows.append(row)

    print(f"\nMatched {matched} / {len(events)} events")

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    tag = f" - {args.label}" if args.label else ""
    out_path = os.path.join(OUTPUT_DIR, f"Benzinga Earnings Beat-Miss{tag}.xlsx")
    try:
        pd.DataFrame(out_rows).to_excel(out_path, index=False)
        print(f"Wrote {out_path}")
    except PermissionError:
        print(f"Could not write {out_path} (likely open in Excel)")


if __name__ == "__main__":
    main()
