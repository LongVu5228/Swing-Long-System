"""
Second pass over All EP Scan.xlsx: for rows where Gap Day Open/Close are still N/A
(and the event date is within the plan's reachable history, i.e. not the 2021-08-10
cutoff case), look up the ticker's historical symbol via Polygon's ticker-events
endpoint and re-pull that day's OHLC under the symbol that was actually active then.

Example: META didn't exist as a ticker until 2022-06-09 (it was FB before that), so
a scan row tagging a pre-rename event as "META" returns nothing until re-queried as "FB".

Only touches the two Gap Day Open/Close cells for rows that are still N/A -- everything
else in the file is left alone.
"""

import argparse
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date

import openpyxl
import requests
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from tqdm import tqdm

load_dotenv()

POLYGON_API_KEY = os.environ.get("POLYGON_API_KEY")
if not POLYGON_API_KEY:
    sys.exit("POLYGON_API_KEY not found. Check your .env file.")

INPUT_XLSX = os.path.join("Files", "EP", "All EP Scan.xlsx")
SHEET_NAME = "OG No Dupes"
PLAN_CUTOFF = date(2021, 8, 10)

BASE_URL = "https://api.polygon.io"
MAX_RETRIES = 5
MAX_WORKERS = 10

SESSION = requests.Session()
_adapter = HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS)
SESSION.mount("https://", _adapter)


def api_get(path, params=None):
    params = dict(params or {})
    params["apiKey"] = POLYGON_API_KEY
    url = f"{BASE_URL}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        resp = SESSION.get(url, params=params, timeout=30)
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code in (403, 404):
            return None
        resp.raise_for_status()
        return resp.json()
    return None


def clean_ticker(raw_ticker):
    t = raw_ticker.split("^")[0].strip()
    if "." in t:
        base, suffix = t.rsplit(".", 1)
        if suffix.isdigit():
            t = base
    return t


def get_ticker_events(ticker):
    data = api_get(f"/vX/reference/tickers/{ticker}/events")
    if not data or "results" not in data:
        return None
    return data["results"].get("events", [])


def resolve_historical_ticker(ticker, event_date):
    """Return the symbol actually in use on event_date, per Polygon's ticker-change history."""
    events = get_ticker_events(ticker)
    lookup_ticker = ticker

    if events is None and ticker.endswith("Q") and len(ticker) > 1:
        # OTC post-bankruptcy "Q" suffix -- try the underlying pre-bankruptcy symbol
        base = ticker[:-1]
        events = get_ticker_events(base)
        if events is not None:
            lookup_ticker = base

    if not events:
        return lookup_ticker

    changes = sorted(
        (e["date"], e["ticker_change"]["ticker"])
        for e in events if e.get("type") == "ticker_change"
    )
    candidate = changes[0][1]
    for d, sym in changes:
        if date.fromisoformat(d) <= event_date:
            candidate = sym
        else:
            break
    return candidate


def get_day_ohlc(ticker, event_date):
    path = f"/v2/aggs/ticker/{ticker}/range/1/day/{event_date.isoformat()}/{event_date.isoformat()}"
    data = api_get(path, {"adjusted": "true"})
    if not data or not data.get("results"):
        return None, None
    bar = data["results"][0]
    return bar["o"], bar["c"]


def resolve_row(ticker_raw, event_date):
    ticker = clean_ticker(ticker_raw)
    resolved = resolve_historical_ticker(ticker, event_date)
    o, c = get_day_ohlc(resolved, event_date)
    return resolved, o, c


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="Print results, don't write to the Excel file")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}

    targets = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[idx["Gap Day Open"]] != "N/A":
            continue
        event_date = date.fromisoformat(str(row[idx["Full Date"]]))
        if event_date < PLAN_CUTOFF:
            continue  # unreachable regardless of symbol -- not what we're fixing here
        targets.append({"row": row_num, "ticker_raw": row[idx["Ticker"]], "event_date": event_date})

    print(f"{len(targets)} N/A rows eligible for rename resolution")

    resolved_cache = {}
    results = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {
            executor.submit(resolve_row, t["ticker_raw"], t["event_date"]): t
            for t in targets
        }
        with tqdm(total=len(futures), desc="Resolving renamed tickers", unit="row") as pbar:
            for future in as_completed(futures):
                t = futures[future]
                resolved, o, c = future.result()
                results[t["row"]] = (t["ticker_raw"], resolved, o, c)
                pbar.update(1)

    recovered = sum(1 for _, _, o, _ in results.values() if o is not None)
    print(f"\nRecovered {recovered} / {len(targets)} rows")

    changed = [(row, orig, resolved) for row, (orig, resolved, o, c) in results.items()
               if o is not None and clean_ticker(orig) != resolved]
    print(f"Symbol changes found: {len(changed)}")
    for row, orig, resolved in changed[:30]:
        print(f"  row {row}: {orig} -> {resolved}")

    if args.dry_run:
        return

    wb2 = openpyxl.load_workbook(INPUT_XLSX, data_only=False)
    ws2 = wb2[SHEET_NAME]
    open_col = header.index("Gap Day Open") + 1
    close_col = header.index("Gap Day Close") + 1

    for row_num, (orig, resolved, o, c) in results.items():
        if o is None:
            continue
        ws2.cell(row=row_num, column=open_col, value=o)
        ws2.cell(row=row_num, column=close_col, value=c)

    wb2.save(INPUT_XLSX)
    print(f"\nWrote {recovered} recovered rows to {INPUT_XLSX}")


if __name__ == "__main__":
    main()
