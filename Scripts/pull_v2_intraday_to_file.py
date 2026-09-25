"""
Same data pull as build_v2_intraday.py (IPO date, pre-gap volume baselines, gap-day
totals, opening-day minute-candle stats) but reads row order from the canonical
"Master Order of Tickers.xlsx" and writes results to a fresh standalone workbook in
Data Pulls/, instead of writing into EP Scan V2.xlsx directly. EP Scan V2.xlsx now
has pivot tables that openpyxl can't safely round-trip on save (that's what caused
the "needs repair" corruption) -- so this file is meant to be copy/pasted in by hand,
row-aligned to the master order.
"""

import argparse
import os
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime

import openpyxl
import pandas as pd
from openpyxl.utils.datetime import from_excel
from tqdm import tqdm

sys.path.insert(0, os.path.dirname(__file__))
from build_v2_intraday import process_ticker, CANDLE_WINDOWS, na, clean_ticker, MAX_WORKERS

MASTER_ORDER_XLSX = os.path.join("Files", "EP", "Master Order of tickers", "Master Order of Tickers.xlsx")
OUTPUT_DIR = os.path.join("Files", "EP", "Data Pulls")


def load_master_order(limit=None):
    wb = openpyxl.load_workbook(MASTER_ORDER_XLSX, data_only=True)
    ws = wb["Sheet1"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ticker_raw = row[idx["Ticker"]]
        fd = row[idx["Full Date"]]
        if not ticker_raw or fd is None:
            continue
        if isinstance(fd, datetime):
            event_date = fd.date()
        elif isinstance(fd, (int, float)):
            event_date = from_excel(fd).date()
        else:
            event_date = date.fromisoformat(str(fd))
        events.append({"row": row_num, "ticker_raw": ticker_raw, "ticker": clean_ticker(ticker_raw),
                        "event_date": event_date})
    if limit:
        events = events[:limit]
    return events


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None)
    args = parser.parse_args()

    events = load_master_order(limit=args.limit)
    by_ticker = {}
    for e in events:
        by_ticker.setdefault(e["ticker"], []).append(e)
    print(f"{len(events)} event rows across {len(by_ticker)} unique tickers (master order)")

    all_results = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {executor.submit(process_ticker, t, evs): t for t, evs in by_ticker.items()}
        with tqdm(total=len(futures), desc="Pulling intraday data", unit="ticker") as pbar:
            for future in as_completed(futures):
                all_results.update(future.result())
                pbar.update(1)

    columns = ["Full Date", "Ticker", "IPO Date",
               "Pre-Gap 30D Avg Dollar Volume", "Pre-Gap 30D Avg Share Volume",
               "Pre-Gap 100D Avg Dollar Volume", "Pre-Gap 100D Avg Share Volume",
               "Gap Day Total Volume", "Gap Day Total Dollar Volume"]
    for m in CANDLE_WINDOWS:
        columns += [f"{m}M Candle Close", f"{m}M Candle Volume", f"{m}M Candle Dollar Volume"]

    out_rows = []
    for e in events:
        r = all_results.get(e["row"], {})
        row = {
            "Full Date": e["event_date"].isoformat(), "Ticker": e["ticker_raw"],
            "IPO Date": na(r.get("ipo_date")),
            "Pre-Gap 30D Avg Dollar Volume": na(r.get("dvol30")),
            "Pre-Gap 30D Avg Share Volume": na(r.get("svol30")),
            "Pre-Gap 100D Avg Dollar Volume": na(r.get("dvol100")),
            "Pre-Gap 100D Avg Share Volume": na(r.get("svol100")),
            "Gap Day Total Volume": na(r.get("gap_vol")),
            "Gap Day Total Dollar Volume": na(r.get("gap_dvol")),
        }
        for m in CANDLE_WINDOWS:
            row[f"{m}M Candle Close"] = na(r.get(f"close_{m}m"))
            row[f"{m}M Candle Volume"] = na(r.get(f"vol_{m}m"))
            row[f"{m}M Candle Dollar Volume"] = na(r.get(f"dvol_{m}m"))
        out_rows.append(row)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    out_path = os.path.join(OUTPUT_DIR, f"V2 Intraday Data - {date.today().isoformat()}.xlsx")
    pd.DataFrame(out_rows, columns=columns).to_excel(out_path, index=False)
    print(f"\nWrote {len(out_rows)} rows to {out_path}")


if __name__ == "__main__":
    main()
