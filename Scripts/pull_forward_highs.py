"""
For every ticker/event-date row in All EP Scan.xlsx, find the highest daily High
reached in the 1/3/6 calendar-month window starting on the event date (inclusive),
and the date that high occurred. Writes the results back as new columns in the
"OG No Dupes" sheet so % gain can be added manually next to each.
"""

import argparse
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone

import openpyxl
import pandas as pd
import requests
from dateutil.relativedelta import relativedelta
from dotenv import load_dotenv
from requests.adapters import HTTPAdapter
from tqdm import tqdm

load_dotenv()

POLYGON_API_KEY = os.environ.get("POLYGON_API_KEY")
if not POLYGON_API_KEY:
    sys.exit("POLYGON_API_KEY not found. Check your .env file.")

INPUT_XLSX = os.path.join("Files", "EP", "All EP Scan.xlsx")
SHEET_NAME = "OG No Dupes"

BASE_URL = "https://api.polygon.io"
MAX_RETRIES = 5
MAX_WORKERS = 20
WINDOWS_MONTHS = [1, 3, 6]

SESSION = requests.Session()
_adapter = HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS)
SESSION.mount("https://", _adapter)

TODAY = date.today()
PLAN_CUTOFF = TODAY - relativedelta(years=5) + relativedelta(days=1)  # this Polygon plan only has 5y of history


def api_get(path, params=None):
    params = dict(params or {})
    params["apiKey"] = POLYGON_API_KEY
    url = f"{BASE_URL}{path}"
    for attempt in range(1, MAX_RETRIES + 1):
        resp = SESSION.get(url, params=params, timeout=30)
        if resp.status_code == 429:
            time.sleep(2 ** attempt)
            continue
        if resp.status_code in (403, 404):
            return None
        resp.raise_for_status()
        return resp.json()
    return None


def clean_ticker(raw_ticker):
    t = raw_ticker.split("^")[0].strip()
    if "." in t:
        base, suffix = t.rsplit(".", 1)
        if suffix.isdigit():  # P123 dedup suffix (ticker reused by a different company), not a share class
            t = base
    return t


def get_daily_bars(ticker, from_date, to_date):
    path = f"/v2/aggs/ticker/{ticker}/range/1/day/{from_date}/{to_date}"
    data = api_get(path, {"adjusted": "true", "sort": "asc", "limit": 50000})
    if data is None or not data.get("results"):
        return None
    df = pd.DataFrame(data["results"])
    df["date"] = df["t"].apply(lambda ms: datetime.fromtimestamp(ms / 1000, tz=timezone.utc).date())
    return df[["date", "h"]].rename(columns={"h": "high"})


def load_events(limit=None):
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    events = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        ticker_raw = row[idx["Ticker"]]
        event_date = row[idx["Full Date"]]
        if not ticker_raw or not event_date:
            continue
        events.append({
            "row": row_num,
            "ticker_raw": ticker_raw,
            "ticker": clean_ticker(ticker_raw),
            "event_date": date.fromisoformat(str(event_date)),
        })
    if limit:
        events = events[:limit]
    return events


def process_ticker_events(ticker, ticker_events):
    lo = min(e["event_date"] for e in ticker_events)
    hi = max(e["event_date"] + relativedelta(months=max(WINDOWS_MONTHS)) for e in ticker_events)
    hi = min(hi, TODAY)
    lo = max(lo, PLAN_CUTOFF)  # clip to what the plan can actually return, rather than 403 the whole request

    bars = get_daily_bars(ticker, lo.isoformat(), hi.isoformat()) if lo <= hi else None

    results = {}
    for e in ticker_events:
        row_result = {}
        truncated = e["event_date"] < PLAN_CUTOFF  # window start predates the plan's 5y history -> high may be understated
        for m in WINDOWS_MONTHS:
            window_end = e["event_date"] + relativedelta(months=m)
            complete = window_end <= TODAY
            if bars is None:
                row_result[m] = (None, None, complete, truncated)
                continue
            mask = (bars["date"] >= e["event_date"]) & (bars["date"] <= window_end)
            window_bars = bars[mask]
            if window_bars.empty:
                row_result[m] = (None, None, complete, truncated)
                continue
            peak_idx = window_bars["high"].idxmax()
            peak_row = window_bars.loc[peak_idx]
            row_result[m] = (peak_row["high"], peak_row["date"].isoformat(), complete, truncated)
        results[e["row"]] = row_result
    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None, help="Only process first N event rows (testing)")
    parser.add_argument("--dry-run", action="store_true", help="Print results, don't write to the Excel file")
    args = parser.parse_args()

    events = load_events(limit=args.limit)
    by_ticker = {}
    for e in events:
        by_ticker.setdefault(e["ticker"], []).append(e)

    print(f"{len(events)} event rows across {len(by_ticker)} unique tickers")

    all_results = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = {
            executor.submit(process_ticker_events, ticker, ticker_events): ticker
            for ticker, ticker_events in by_ticker.items()
        }
        with tqdm(total=len(futures), desc="Pulling forward highs", unit="ticker") as pbar:
            for future in as_completed(futures):
                all_results.update(future.result())
                pbar.update(1)

    if args.dry_run:
        preview_rows = []
        for e in events:
            r = all_results.get(e["row"], {})
            preview_rows.append({
                "ticker": e["ticker"],
                "event_date": e["event_date"].isoformat(),
                **{f"high_{m}m": r.get(m, (None, None, None, None))[0] for m in WINDOWS_MONTHS},
                **{f"high_{m}m_date": r.get(m, (None, None, None, None))[1] for m in WINDOWS_MONTHS},
                **{f"{m}m_complete": r.get(m, (None, None, None, None))[2] for m in WINDOWS_MONTHS},
                **{f"{m}m_plan_truncated": r.get(m, (None, None, None, None))[3] for m in WINDOWS_MONTHS},
            })
        print(pd.DataFrame(preview_rows).to_string(index=False))
        return

    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=False)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    next_col = len(header) + 1

    col_map = {}
    for m in WINDOWS_MONTHS:
        for suffix in ["High", "High Date", "Window Complete", "Plan Data Truncated"]:
            ws.cell(row=1, column=next_col, value=f"{m}M {suffix}")
            col_map[(m, suffix)] = next_col
            next_col += 1

    for e in events:
        r = all_results.get(e["row"])
        if not r:
            continue
        for m in WINDOWS_MONTHS:
            high, high_date, complete, truncated = r[m]
            ws.cell(row=e["row"], column=col_map[(m, "High")], value=high if high is not None else "N/A")
            ws.cell(row=e["row"], column=col_map[(m, "High Date")], value=high_date if high_date is not None else "N/A")
            ws.cell(row=e["row"], column=col_map[(m, "Window Complete")], value=complete)
            ws.cell(row=e["row"], column=col_map[(m, "Plan Data Truncated")], value=truncated)

    wb.save(INPUT_XLSX)
    print(f"\nWrote {len(events)} rows of forward-high data to {INPUT_XLSX}")


if __name__ == "__main__":
    main()
